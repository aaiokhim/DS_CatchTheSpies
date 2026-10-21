import sys
from pathlib import Path
from openpyxl import load_workbook
import csv
import re

sys.path.insert(
    0,
    str(
        next(
            p for p in Path(__file__).resolve().parents
            if (p / "paths.py").exists()
        )
    )
)

from paths import RAW, P3

folder = RAW / "YourBoardingPassDotAero"

output_file = P3 / "output" / "boarding_passes.csv"

errors_file = P3 / "intermediate" / "boarding_passes_errors.csv"



def clean_value(value):

    if value is None:
        return None

    if isinstance(value, str):
        value = value.strip()

        if value == "":
            return None

    return value


def parse_loyalty(value):

    value = clean_value(value)

    if value is None:
        return None, None, None

    raw_value = str(value)

    match = re.fullmatch(
        r"([A-Z0-9]{2})\s+(\d+)",
        raw_value
    )

    if match:
        return (
            raw_value,
            match.group(1),
            match.group(2)
        )

    return raw_value, None, None


def is_valid_boarding_pass(ws):

    return (
        clean_value(ws["A1"].value) == "BOARDING PASS"
        and clean_value(ws["F1"].value) == "SEQUENCE:"
        and clean_value(ws["A7"].value) == "GATE"
        and clean_value(ws["G11"].value) == "SEAT"
        and clean_value(ws["A13"].value) == "PNR:"
        and clean_value(ws["D13"].value) == "E-TICKET"
    )


def get_date_from_filename(filename):

    match = re.search(
        r"(\d{4}-\d{2}-\d{2})",
        filename
    )

    if match:
        return match.group(1)

    return None


def parse_sheet(ws, source_file):

    loyalty_raw, loyalty_code, loyalty_number = parse_loyalty(
        ws["F3"].value
    )

    source_date = get_date_from_filename(source_file.name)

    flight_date = clean_value(ws["A9"].value)

    if flight_date is not None:
        flight_date = str(flight_date)

    if source_date is not None and flight_date is not None:
        date_matches_filename = source_date == flight_date
    else:
        date_matches_filename = None

    return {
        "title": clean_value(ws["A3"].value),

        "passenger_name_raw":
            clean_value(ws["B3"].value),

        "loyalty_id_raw":
            loyalty_raw,

        "loyalty_program_code":
            loyalty_code,

        "loyalty_number":
            loyalty_number,

        "booking_class_raw":
            clean_value(ws["H3"].value),

        "sequence":
            clean_value(ws["H1"].value),

        "flight_number":
            clean_value(ws["A5"].value),

        "origin_city":
            clean_value(ws["D5"].value),

        "destination_city":
            clean_value(ws["H5"].value),

        "gate":
            clean_value(ws["B7"].value),

        "origin_airport":
            clean_value(ws["D7"].value),

        "destination_airport":
            clean_value(ws["H7"].value),

        "flight_date":
            flight_date,

        "departure_time":
            clean_value(ws["C9"].value),

        "operated_by_raw":
            clean_value(ws["E9"].value),

        "seat":
            clean_value(ws["H11"].value),

        "pnr":
            clean_value(ws["B13"].value),

        "e_ticket":
            clean_value(ws["E13"].value),

        "source_file":
            source_file.name,

        "source_sheet":
            ws.title,

        "source_date":
            source_date,

        "date_matches_filename":
            date_matches_filename,
    }



fieldnames = [
    "title",
    "passenger_name_raw",

    "loyalty_id_raw",
    "loyalty_program_code",
    "loyalty_number",

    "booking_class_raw",
    "sequence",

    "flight_number",

    "origin_city",
    "destination_city",

    "gate",

    "origin_airport",
    "destination_airport",

    "flight_date",
    "departure_time",

    "operated_by_raw",

    "seat",

    "pnr",
    "e_ticket",

    "source_file",
    "source_sheet",
    "source_date",
    "date_matches_filename",
]


files = sorted(folder.glob("*.xlsx"))

print("Найдено XLSX:", len(files))

if len(files) == 0:
    print("Файлы не найдены.")
    print("Проверь путь:")
    print(folder)
    raise SystemExit


total_sheets = 0
total_parsed = 0
total_invalid = 0
total_errors = 0


with open(
    output_file,
    "w",
    newline="",
    encoding="utf-8-sig"
) as csvfile, open(
    errors_file,
    "w",
    newline="",
    encoding="utf-8-sig"
) as errorfile:

    writer = csv.DictWriter(
        csvfile,
        fieldnames=fieldnames
    )

    writer.writeheader()

    error_writer = csv.DictWriter(
        errorfile,
        fieldnames=[
            "source_file",
            "source_sheet",
            "error_type",
            "message",
        ]
    )

    error_writer.writeheader()


    for file_number, file in enumerate(files, start=1):

        print(
            f"[{file_number}/{len(files)}] "
            f"Обрабатывается: {file.name}"
        )

        try:
            wb = load_workbook(
                file,
                read_only=True,
                data_only=True
            )

        except Exception as e:
            total_errors += 1

            error_writer.writerow({
                "source_file": file.name,
                "source_sheet": "",
                "error_type": "workbook_open_error",
                "message": str(e),
            })

            continue

        for ws in wb.worksheets:

            total_sheets += 1

            try:
                if not is_valid_boarding_pass(ws):

                    total_invalid += 1

                    error_writer.writerow({
                        "source_file": file.name,
                        "source_sheet": ws.title,
                        "error_type": "unexpected_template",
                        "message":
                            "Лист не соответствует ожидаемому шаблону",
                    })

                    continue

                row = parse_sheet(
                    ws,
                    file
                )

                writer.writerow(row)

                total_parsed += 1

            except Exception as e:

                total_errors += 1

                error_writer.writerow({
                    "source_file": file.name,
                    "source_sheet": ws.title,
                    "error_type": "sheet_parse_error",
                    "message": str(e),
                })

        wb.close()

print(
    f"Готово: {total_parsed} листов спарсено, "
    f"{total_invalid} неправильных шаблонов, "
    f"{total_errors} ошибок."
)
